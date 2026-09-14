import hashlib
import json
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

APP_ROOT = Path(__file__).parent / "portable" / "app"
sys.path.insert(0, str(APP_ROOT))

from process_list import (  # noqa: E402
    INPUT_MANIFEST_SCHEMA_VERSION,
    REPORT_HEADERS,
    import_process_workbook,
    _positive_integer,
    validate_input_manifest,
    write_analysis_report,
)


def make_workbook(path: Path, rows, headers=("numero_processo", "ano_processo")):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Planilha2"
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)


class ProcessListImportTests(unittest.TestCase):
    def test_imports_valid_rows_and_freezes_order_and_duplicates(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "entrada.xlsx"
            make_workbook(path, [(101, 2023), (202, 2024), (101, 2023), (303, 2025)])

            manifest = import_process_workbook(path)

            self.assertEqual(manifest["schema_version"], INPUT_MANIFEST_SCHEMA_VERSION)
            self.assertEqual(manifest["row_count"], 4)
            self.assertEqual(manifest["unique_count"], 3)
            self.assertEqual(manifest["duplicate_count"], 1)
            self.assertEqual(
                manifest["ordered_unique_keys"], ["101/2023", "202/2024", "303/2025"]
            )
            self.assertEqual(manifest["rows"][0]["source_row"], 2)
            self.assertIsNone(manifest["rows"][0]["duplicate_of_row"])
            self.assertEqual(manifest["rows"][2]["duplicate_of_row"], 2)
            self.assertRegex(manifest["input_list_id"], r"^input-[0-9a-f]{24}$")
            self.assertEqual(
                manifest["input_sha256"], hashlib.sha256(path.read_bytes()).hexdigest()
            )

    def test_rejects_missing_headers_empty_cells_and_non_integer_values(self):
        cases = [
            (("numero_processo", "wrong"), [(1, 2023)], "cabeçalho"),
            (("numero_processo", "ano_processo"), [(None, 2023)], "vazia"),
            (("numero_processo", "ano_processo"), [(1.5, 2023)], "inteiro"),
            (("numero_processo", "ano_processo"), [(1, "2023")], "inteiro"),
            (("numero_processo", "ano_processo"), [(0, 2023)], "positivo"),
        ]
        for index, (headers, rows, expected) in enumerate(cases):
            with self.subTest(index=index):
                with TemporaryDirectory() as directory:
                    path = Path(directory) / "entrada.xlsx"
                    make_workbook(path, rows, headers=headers)
                    with self.assertRaisesRegex(ValueError, expected):
                        import_process_workbook(path)

    def test_accepts_excel_integral_numbers_serialized_as_float(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "entrada.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Planilha2"
            sheet.append(["numero_processo", "ano_processo"])
            sheet.append([101.0, 2023.0])
            sheet["A2"].number_format = "0.0"
            sheet["B2"].number_format = "0.0"
            workbook.save(path)

            manifest = import_process_workbook(path)

            self.assertEqual(manifest["ordered_unique_keys"], ["101/2023"])

    def test_accepts_only_integral_float_values_as_excel_integers(self):
        self.assertEqual(_positive_integer(101.0, "numero_processo", 2), 101)
        with self.assertRaisesRegex(ValueError, "inteiro"):
            _positive_integer(101.5, "numero_processo", 2)

    def test_report_has_one_row_per_input_line_and_reconciled_summary(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "entrada.xlsx"
            output = Path(directory) / "resultado.xlsx"
            make_workbook(path, [(101, 2023), (101, 2023), (202, 2024)])
            manifest = import_process_workbook(path)
            classifications = {
                "101/2023": {
                    "marker_present": True,
                    "marker_observed": "PROFESSOR - IPERN",
                    "area_status": "PRECISA_COMPLEMENTAR",
                    "action_signature": {
                        "alt": "Complementar Ato",
                        "title": "Complementar Ato",
                        "src": "red-icon.png",
                    },
                    "lot_number": 1,
                    "econtas_status": "downloaded",
                    "documents_downloaded": 2,
                    "error": None,
                },
                "202/2024": {
                    "marker_present": False,
                    "marker_observed": None,
                    "area_status": "ATO_COMPLEMENTADO",
                    "action_signature": None,
                    "lot_number": None,
                    "econtas_status": "not_started",
                    "documents_downloaded": 0,
                    "error": None,
                },
            }

            result = write_analysis_report(output, manifest, classifications)

            self.assertEqual(result["row_count"], 3)
            self.assertEqual(result["unique_count"], 2)
            self.assertEqual(result["duplicate_count"], 1)
            self.assertEqual(result["eligible_count"], 1)
            self.assertTrue(output.exists())
            workbook = load_workbook(output, read_only=True, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["Resultados", "Resumo"])
                rows = list(workbook["Resultados"].iter_rows(values_only=True))
                self.assertEqual(tuple(rows[0]), REPORT_HEADERS)
                self.assertEqual(len(rows) - 1, 3)
                self.assertEqual(rows[2][2], "101/2023")
                self.assertEqual(rows[2][3], 2)
            finally:
                workbook.close()

    def test_manifest_validation_rejects_tampered_identity_and_row_provenance(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "entrada.xlsx"
            make_workbook(path, [(101, 2023), (101, 2023), (202, 2024)])
            manifest = import_process_workbook(path)

            tampered_hash = dict(manifest, input_sha256="z" * 64)
            with self.assertRaisesRegex(ValueError, "input_sha256"):
                validate_input_manifest(tampered_hash)

            tampered_row = json.loads(json.dumps(manifest))
            tampered_row["rows"][1]["process_key"] = "999/2024"
            with self.assertRaisesRegex(ValueError, "ordem|chaves|linha"):
                validate_input_manifest(tampered_row)


if __name__ == "__main__":
    unittest.main()
