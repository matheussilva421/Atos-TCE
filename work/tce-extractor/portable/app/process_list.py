"""Import and reporting contracts for the authoritative process workbook.

The workbook is treated as an input artifact, never edited in place.  This
module intentionally keeps absolute paths out of persisted JSON so a manifest
can be moved between the menu, bridge and extension without leaking local
filesystem details.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
import base64
import hashlib
import json
import math
import os
from pathlib import Path
import re
from tempfile import NamedTemporaryFile
from typing import Any


INPUT_MANIFEST_SCHEMA_VERSION = 1
REPORT_HEADERS = (
    "linha_original",
    "numero_processo",
    "processo",
    "referencia_duplicidade",
    "presente_no_marcador",
    "marcador_observado",
    "classificacao_area_restrita",
    "assinatura_controle_alt",
    "assinatura_controle_title",
    "assinatura_controle_src",
    "numero_lote",
    "estado_econtas",
    "documentos_baixados",
    "erro",
)
SUMMARY_FIELDS = (
    "linhas",
    "unicos",
    "duplicados",
    "elegiveis",
    "complementados",
    "ausentes",
    "ambiguos",
    "bloqueados",
    "lotes",
)
_REQUIRED_HEADERS = ("numero_processo", "ano_processo")
_INPUT_LIST_ID_RE = re.compile(r"^input-[0-9a-f]{24}$")
_SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
_PROCESS_KEY_RE = re.compile(r"^[1-9][0-9]*/[1-9][0-9]{3,}$")


def _workbook_module():
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - exercised by runtime gate
        raise RuntimeError(
            "openpyxl não está disponível no runtime portátil; execute o construtor do pacote"
        ) from error
    return load_workbook


def _positive_integer(value: object, field: str, row: int) -> int:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"{field} da linha {row} não pode ser vazia")
    if type(value) is float:
        if not math.isfinite(value) or not value.is_integer():
            raise ValueError(f"{field} da linha {row} deve ser inteiro")
        value = int(value)
    elif type(value) is not int:
        raise ValueError(f"{field} da linha {row} deve ser inteiro")
    if value <= 0:
        raise ValueError(f"{field} da linha {row} deve ser positivo")
    return value


def _canonical_key(number: object, year: object, row: int) -> tuple[int, int, str]:
    normalized_number = _positive_integer(number, "numero_processo", row)
    normalized_year = _positive_integer(year, "ano_processo", row)
    return normalized_number, normalized_year, f"{normalized_number}/{normalized_year}"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def import_process_workbook(path: str | Path, *, sheet_name: str = "Planilha2") -> dict[str, Any]:
    """Read and validate a process workbook into an immutable JSON manifest."""

    source = Path(path)
    if source.suffix.casefold() != ".xlsx" or not source.is_file():
        raise ValueError("arquivo de entrada deve ser um .xlsx existente")
    source_hash = _sha256(source)
    load_workbook = _workbook_module()
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as error:
        raise ValueError(f"não foi possível ler o .xlsx: {error}") from error
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"planilha ausente: {sheet_name}")
        worksheet = workbook[sheet_name]
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        raise ValueError("planilha sem cabeçalho")

    header = [str(value).strip() if value is not None else "" for value in rows[0]]
    positions: dict[str, int] = {}
    for required in _REQUIRED_HEADERS:
        if required not in header:
            raise ValueError(f"cabeçalho obrigatório ausente: {required}")
        position = header.index(required)
        if required in positions:
            raise ValueError(f"cabeçalho duplicado: {required}")
        positions[required] = position

    data_rows = list(rows[1:])
    while data_rows and all(value is None for value in data_rows[-1]):
        data_rows.pop()
    if not data_rows:
        raise ValueError("planilha sem processos")

    manifest_rows: list[dict[str, Any]] = []
    first_rows: dict[str, int] = {}
    ordered_unique_keys: list[str] = []
    for offset, values in enumerate(data_rows, start=2):
        number = values[positions["numero_processo"]] if positions["numero_processo"] < len(values) else None
        year = values[positions["ano_processo"]] if positions["ano_processo"] < len(values) else None
        number, year, process_key = _canonical_key(number, year, offset)
        duplicate_of = first_rows.get(process_key)
        if duplicate_of is None:
            first_rows[process_key] = offset
            ordered_unique_keys.append(process_key)
        manifest_rows.append(
            {
                "source_row": offset,
                "numero_processo": number,
                "ano_processo": year,
                "process_key": process_key,
                "duplicate_of_row": duplicate_of,
            }
        )

    unique_count = len(ordered_unique_keys)
    return {
        "schema_version": INPUT_MANIFEST_SCHEMA_VERSION,
        "input_list_id": f"input-{source_hash[:24]}",
        "input_sha256": source_hash,
        "source_filename": source.name,
        "sheet_name": sheet_name,
        "headers": header,
        "row_count": len(manifest_rows),
        "unique_count": unique_count,
        "duplicate_count": len(manifest_rows) - unique_count,
        "ordered_unique_keys": ordered_unique_keys,
        "rows": manifest_rows,
    }


def _classification_for(classifications: Mapping[str, Mapping[str, Any]], key: str) -> Mapping[str, Any]:
    value = classifications.get(key, {})
    if not isinstance(value, Mapping):
        return {}
    return value


def _summary(manifest: Mapping[str, Any], classifications: Mapping[str, Mapping[str, Any]]) -> dict[str, int]:
    values = [_classification_for(classifications, key) for key in manifest["ordered_unique_keys"]]
    statuses = [str(value.get("area_status", "BLOQUEADO")) for value in values]
    eligible = sum(status == "PRECISA_COMPLEMENTAR" for status in statuses)
    completed = sum(status == "ATO_COMPLEMENTADO" for status in statuses)
    absent = sum(status == "NAO_ENCONTRADO_AREA_RESTRITA" for status in statuses)
    ambiguous = sum(status == "AMBIGUO" for status in statuses)
    blocked = sum(status == "BLOQUEADO" for status in statuses)
    lots = {
        value.get("lot_number")
        for value in values
        if isinstance(value.get("lot_number"), int) and value["lot_number"] > 0
    }
    return {
        "linhas": int(manifest["row_count"]),
        "unicos": int(manifest["unique_count"]),
        "duplicados": int(manifest["duplicate_count"]),
        "elegiveis": eligible,
        "complementados": completed,
        "ausentes": absent,
        "ambiguos": ambiguous,
        "bloqueados": blocked,
        "lotes": len(lots),
    }


def write_analysis_report(
    output_path: str | Path,
    manifest: Mapping[str, Any],
    classifications: Mapping[str, Mapping[str, Any]],
) -> dict[str, int]:
    """Write a derived audit workbook and return its reconciled totals."""

    from openpyxl import Workbook

    output = Path(output_path)
    if output.suffix.casefold() != ".xlsx":
        raise ValueError("relatório deve ser um .xlsx")
    rows = manifest.get("rows")
    if not isinstance(rows, Sequence):
        raise ValueError("manifesto de entrada sem linhas")
    summary = _summary(manifest, classifications)
    workbook = Workbook()
    results = workbook.active
    results.title = "Resultados"
    results.append(list(REPORT_HEADERS))
    for entry in rows:
        key = entry["process_key"]
        value = _classification_for(classifications, key)
        signature = value.get("action_signature")
        if not isinstance(signature, Mapping):
            signature = {}
        duplicate = entry.get("duplicate_of_row")
        results.append(
            [
                entry["source_row"],
                entry["numero_processo"],
                key,
                duplicate,
                value.get("marker_present", False),
                value.get("marker_observed"),
                value.get("area_status", "BLOQUEADO"),
                signature.get("alt"),
                signature.get("title"),
                signature.get("src"),
                value.get("lot_number"),
                value.get("econtas_status", "not_started"),
                value.get("documents_downloaded", 0),
                value.get("error"),
            ]
        )
    summary_sheet = workbook.create_sheet("Resumo")
    summary_sheet.append(["metrica", "valor"])
    for field in SUMMARY_FIELDS:
        summary_sheet.append([field, summary[field]])
    summary_sheet.append(["input_list_id", manifest.get("input_list_id")])
    summary_sheet.append(["input_sha256", manifest.get("input_sha256")])
    summary_sheet.append(["sheet_name", manifest.get("sheet_name")])

    output.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(prefix=f".{output.stem}-", suffix=".tmp.xlsx", dir=output.parent, delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        temporary_path.replace(output)
    finally:
        temporary_path.unlink(missing_ok=True)
    return {
        **summary,
        "row_count": summary["linhas"],
        "unique_count": summary["unicos"],
        "duplicate_count": summary["duplicados"],
        "eligible_count": summary["elegiveis"],
        "lot_count": summary["lotes"],
    }


def validate_input_manifest(value: Mapping[str, Any]) -> dict[str, Any]:
    """Validate a persisted manifest without needing the original workbook."""

    if not isinstance(value, Mapping) or value.get("schema_version") != INPUT_MANIFEST_SCHEMA_VERSION:
        raise ValueError("manifesto da lista inválido")
    for field in ("input_list_id", "input_sha256", "source_filename", "sheet_name"):
        if not isinstance(value.get(field), str) or not value[field].strip():
            raise ValueError(f"manifesto sem {field}")
    if not _INPUT_LIST_ID_RE.fullmatch(value["input_list_id"]):
        raise ValueError("input_list_id inválido")
    if not _SHA256_RE.fullmatch(value["input_sha256"].casefold()):
        raise ValueError("input_sha256 inválido")
    if Path(value["source_filename"]).name != value["source_filename"] or Path(value["source_filename"]).suffix.casefold() != ".xlsx":
        raise ValueError("source_filename inválido")
    headers = value.get("headers")
    if not isinstance(headers, list) or any(not isinstance(header, str) for header in headers):
        raise ValueError("cabeçalhos do manifesto inválidos")
    for required in _REQUIRED_HEADERS:
        if headers.count(required) != 1:
            raise ValueError(f"cabeçalho obrigatório inválido: {required}")
    counts = ("row_count", "unique_count", "duplicate_count")
    for field in counts:
        if type(value.get(field)) is not int or value[field] < 0:
            raise ValueError(f"{field} do manifesto inválido")
    rows = value.get("rows")
    keys = value.get("ordered_unique_keys")
    if not isinstance(rows, list) or not isinstance(keys, list) or not keys:
        raise ValueError("manifesto sem linhas ou chaves únicas")
    if value["row_count"] <= 0 or value["unique_count"] <= 0:
        raise ValueError("totais do manifesto devem ser positivos")
    if value["duplicate_count"] != value["row_count"] - value["unique_count"]:
        raise ValueError("totais de duplicidade do manifesto não conferem")
    if len(rows) != value["row_count"] or len(keys) != value["unique_count"]:
        raise ValueError("totais do manifesto não conferem")
    if any(not isinstance(key, str) or not _PROCESS_KEY_RE.fullmatch(key) for key in keys):
        raise ValueError("chave de processo inválida")
    if len(set(keys)) != len(keys):
        raise ValueError("chaves únicas repetidas no manifesto")
    seen: set[str] = set()
    first_rows: dict[str, int] = {}
    source_rows: set[int] = set()
    for row in rows:
        if not isinstance(row, Mapping) or not isinstance(row.get("process_key"), str):
            raise ValueError("linha do manifesto inválida")
        key = row["process_key"]
        source_row = row.get("source_row")
        number = row.get("numero_processo")
        year = row.get("ano_processo")
        if type(source_row) is not int or source_row < 2 or source_row in source_rows:
            raise ValueError("linha de origem inválida")
        if type(number) is not int or number <= 0 or type(year) is not int or year <= 0:
            raise ValueError("número/ano da linha inválido")
        if key != f"{number}/{year}" or not _PROCESS_KEY_RE.fullmatch(key):
            raise ValueError("chave da linha não é canônica")
        source_rows.add(source_row)
        if row.get("duplicate_of_row") is None:
            if key in seen:
                raise ValueError("primeira ocorrência duplicada no manifesto")
            seen.add(key)
            first_rows[key] = source_row
        else:
            duplicate_of = row["duplicate_of_row"]
            if type(duplicate_of) is not int or duplicate_of != first_rows.get(key):
                raise ValueError("referência de duplicidade inválida")
    if seen != set(keys):
        raise ValueError("ordem de chaves únicas do manifesto não confere")
    if [row["process_key"] for row in rows if row.get("duplicate_of_row") is None] != keys:
        raise ValueError("ordem de chaves únicas do manifesto não confere")
    return dict(value)


def _atomic_json(path: Path, value: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(prefix=f".{path.name}-", suffix=".tmp", dir=path.parent, delete=False, mode="w", encoding="utf-8", newline="\n") as temporary:
        temporary_path = Path(temporary.name)
        json.dump(value, temporary, ensure_ascii=False, indent=2)
        temporary.write("\n")
        temporary.flush()
        os.fsync(temporary.fileno())
    try:
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)


class ProcessListStore:
    """Durable active-list registry shared by menu, bridge and extension."""

    def __init__(self, root: str | Path) -> None:
        self.root = Path(root).resolve()
        self.directory = self.root / "automacao" / "listas"
        self.directory.mkdir(parents=True, exist_ok=True)
        self.active_path = self.directory / "active.json"

    def save(self, manifest: Mapping[str, Any]) -> dict[str, Any]:
        payload = validate_input_manifest(manifest)
        list_id = payload["input_list_id"]
        _atomic_json(self.directory / f"{list_id}.json", payload)
        _atomic_json(self.active_path, {"input_list_id": list_id})
        return dict(payload)

    def load(self, input_list_id: str) -> dict[str, Any]:
        if not isinstance(input_list_id, str) or not input_list_id.startswith("input-"):
            raise ValueError("input_list_id inválido")
        try:
            payload = json.loads((self.directory / f"{input_list_id}.json").read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError) as error:
            raise ValueError("lista de processos não encontrada") from error
        return validate_input_manifest(payload)

    def active(self) -> dict[str, Any] | None:
        if not self.active_path.is_file():
            return None
        try:
            pointer = json.loads(self.active_path.read_text(encoding="utf-8"))
            return self.load(pointer["input_list_id"])
        except (KeyError, OSError, UnicodeError, json.JSONDecodeError, ValueError) as error:
            raise ValueError("lista ativa inválida") from error


def import_process_workbook_bytes(root: str | Path, filename: str, content: str) -> dict[str, Any]:
    """Import a base64-encoded workbook received through the authenticated bridge."""

    if not isinstance(filename, str) or Path(filename).name != filename or Path(filename).suffix.casefold() != ".xlsx":
        raise ValueError("filename deve ser um nome .xlsx simples")
    if not isinstance(content, str) or not content:
        raise ValueError("conteúdo da lista ausente")
    try:
        raw = base64.b64decode(content.encode("ascii"), validate=True)
    except (ValueError, UnicodeEncodeError) as error:
        raise ValueError("conteúdo base64 inválido") from error
    if not raw:
        raise ValueError("conteúdo da lista vazio")
    directory = Path(root).resolve() / "automacao" / "listas"
    directory.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with NamedTemporaryFile(prefix=".incoming-", suffix=".xlsx", dir=directory, delete=False) as temporary:
            temporary.write(raw)
            temporary_path = Path(temporary.name)
        manifest = import_process_workbook(temporary_path)
        manifest["source_filename"] = filename
        return ProcessListStore(root).save(manifest)
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
